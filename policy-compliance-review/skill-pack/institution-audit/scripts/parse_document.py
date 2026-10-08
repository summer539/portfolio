#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""企业制度 DOCX 结构化解析器，支持 Word 自动编号。"""
import argparse, json, os, zipfile
from lxml import etree
from docx import Document

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
DIGITS = "零一二三四五六七八九"

def element_val(el):
    return el.get(f"{W}val") if el is not None else None

def cn_num(n):
    if n < 10: return DIGITS[n]
    if n == 10: return "十"
    if n < 20: return "十" + (DIGITS[n % 10] if n % 10 else "")
    t, o = divmod(n, 10)
    return DIGITS[t] + "十" + (DIGITS[o] if o else "")

def load_numbering(path):
    with zipfile.ZipFile(path) as z:
        if "word/numbering.xml" not in z.namelist(): return {}, {}
        root = etree.fromstring(z.read("word/numbering.xml"))
    nums = {}
    for num in root.findall(f"{W}num"):
        nums[num.get(f"{W}numId")] = element_val(num.find(f"{W}abstractNumId"))
    levels = {}
    for abstract in root.findall(f"{W}abstractNum"):
        aid = abstract.get(f"{W}abstractNumId")
        levels[aid] = {}
        for lvl in abstract.findall(f"{W}lvl"):
            levels[aid][lvl.get(f"{W}ilvl")] = (
                element_val(lvl.find(f"{W}numFmt")),
                element_val(lvl.find(f"{W}lvlText")) or "%1")
    return nums, levels

def parse_explicit(path):
    """兼容编号已写入文本、但同时带有异常 Word numPr 的制度文档。"""
    import re
    doc = Document(path)
    chapters, current = [], None
    headings, started, attachment = [], False, None
    chapter_re = re.compile(r"^(第[零一二两三四五六七八九十百千万]+章)\s*(.*)$")
    article_re = re.compile(r"^(第[零一二两三四五六七八九十百千万]+条)\s*(.*)$")
    attach_re = re.compile(r"^(附件\d+[:：].*)$")
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text: continue
        if attach_re.match(text):
            attachment = text; started = True; continue
        m = chapter_re.match(text)
        if m:
            started = True
            current = {"label": m.group(1), "name": m.group(2).strip(), "articles": []}
            chapters.append(current); continue
        m = article_re.match(text)
        if m:
            started = True
            if current is None:
                current = {"label": "", "name": "正文", "articles": []}; chapters.append(current)
            current["articles"].append({"label": m.group(1), "content": m.group(2).strip()})
            continue
        if not started:
            headings.append(text); continue
        if current is not None and not attachment:
            current["articles"][-1]["content"] += ("\n" if current["articles"][-1]["content"] else "") + text
    title = headings[-1] if headings else os.path.splitext(os.path.basename(path))[0]
    # 附件表格作为结构化数据保留，供审查时核对正文引用
    attachments = []
    for i, table in enumerate(doc.tables, 1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        attachments.append({"table_no": i, "rows": rows})
    return {"title": title, "source": os.path.abspath(path), "chapters": chapters, "attachments": attachments}


def parse(path):
    nums, levels = load_numbering(path)
    doc = Document(path)
    visible = "\n".join(p.text for p in doc.paragraphs)
    # 某些 Word 文档把 numPr 设为异常 ilvl=255，但编号已实际写入文本。
    # 优先使用显式编号解析，避免把正文误判为空。
    if "第一章" in visible and "第一条" in visible:
        return parse_explicit(path)
    counters = {}
    chapters, current = [], None
    headings = []
    started = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text: continue
        ppr = paragraph._p.pPr
        num_id = ilvl = None
        if ppr is not None and ppr.numPr is not None:
            num_id = str(ppr.numPr.numId.val) if ppr.numPr.numId is not None else None
            ilvl = str(ppr.numPr.ilvl.val) if ppr.numPr.ilvl is not None else "0"
        label = ""
        if num_id in nums:
            started = True
            key = (num_id, ilvl)
            counters[key] = counters.get(key, 0) + 1
            aid = nums[num_id]
            fmt, template = levels.get(aid, {}).get(ilvl, ("decimal", "%1"))
            value = cn_num(counters[key]) if fmt in ("chineseCounting", "chineseCountingThousand") else str(counters[key])
            label = template.replace("%1", value).strip()
        elif not started:
            headings.append(text)
            continue
        if "条" in label:
            if current is None:
                current = {"label": "", "name": "正文", "articles": []}
                chapters.append(current)
            current["articles"].append({"label": label, "content": text})
        elif label:
            current = {"label": label, "name": text, "articles": []}
            chapters.append(current)
        else:
            if current is None:
                current = {"label": "", "name": "正文", "articles": []}
                chapters.append(current)
            current["articles"].append({"label": "", "content": text})
    title = headings[-1] if headings else os.path.splitext(os.path.basename(path))[0]
    return {"title": title, "source": os.path.abspath(path), "chapters": chapters}

def to_md(tree):
    lines = [f"# {tree['title']}", ""]
    total = 0
    for chapter in tree["chapters"]:
        lines.append(f"## {(chapter['label'] + ' ' + chapter['name']).strip()}")
        for article in chapter["articles"]:
            total += 1
            lines.append(f"- {(article['label'] + ' ' + article['content']).strip()}")
        lines.append("")
    lines.append(f"---\n共 {len(tree['chapters'])} 章 / {total} 条")
    return "\n".join(lines)

def main():
    p = argparse.ArgumentParser()
    p.add_argument("docx")
    p.add_argument("--json", required=True)
    p.add_argument("--md")
    args = p.parse_args()
    tree = parse(args.docx)
    os.makedirs(os.path.dirname(os.path.abspath(args.json)), exist_ok=True)
    with open(args.json, "w", encoding="utf-8") as f:
        json.dump(tree, f, ensure_ascii=False, indent=2)
    if args.md:
        with open(args.md, "w", encoding="utf-8") as f: f.write(to_md(tree))
    articles = sum(len(c["articles"]) for c in tree["chapters"])
    print(f"结构化完成: {tree['title']} | {len(tree['chapters'])} 章 | {articles} 条")
    print(f"JSON: {args.json}")
    if args.md: print(f"Markdown: {args.md}")

if __name__ == "__main__": main()
