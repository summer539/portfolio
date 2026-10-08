from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from audit_common import (
    ARTICLE_RE, CHAPTER_RE, SECTION_RE,
    QUICK_AUDIT_PROFILE, build_doc_structure, concern_sort_key,
    infer_organization_scope, numbering_index, paragraph_numbering,
    render_auto_label, structure_label,
)

SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_ROOT = SCRIPT_DIR.parent
DEFAULT_OUTPUT_DIR = Path.cwd() / "审核输出" / "normative"
SUFFIXES = ["章程", "员工手册", "管理办法", "管理规定", "实施细则", "操作规程", "指导手册", "工作指引", "议事规则", "工作规则", "工作细则"]
TITLE_SUFFIXES = SUFFIXES
FONT_ALIASES = {"宋体": {"宋体", "SimSun"}, "黑体": {"黑体", "SimHei"}, "仿宋": {"仿宋", "仿宋_GB2312", "FangSong"}}
FORMAT_STANDARD = {
    "chapter": {"font": "黑体", "size_pt": 16.0, "bold": True, "alignment": "居中"},
    "body": {"font": "仿宋", "size_pt": 16.0, "alignment_required": False, "first_indent_required": False},
    "page": {"width_in": 8.27, "height_in": 11.69, "top_cm": 2.54, "bottom_cm": 2.54, "left_cm": 3.17, "right_cm": 3.17},
}


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def xml_font(run) -> list[str]:
    rfonts = run._r.rPr.rFonts if run._r.rPr is not None and run._r.rPr.rFonts is not None else None
    if rfonts is None:
        return []
    return [value for value in [rfonts.get(qn("w:eastAsia")), rfonts.get(qn("w:ascii")), rfonts.get(qn("w:hAnsi"))] if value]


def run_fonts(run, paragraph) -> list[str]:
    direct = [x for x in [run.font.name] if x] + xml_font(run)
    if direct:
        return sorted(set(direct))
    style = paragraph.style
    style_name = style.font.name if style is not None else None
    return [style_name] if style_name else []


def run_size(run, paragraph) -> float | None:
    if run.font.size is not None:
        return round(run.font.size.pt, 2)
    if paragraph.style is not None and paragraph.style.font.size is not None:
        return round(paragraph.style.font.size.pt, 2)
    return None


def run_bold(run, paragraph) -> bool | None:
    if run.bold is not None:
        return run.bold
    if paragraph.style is not None and paragraph.style.font.bold is not None:
        return paragraph.style.font.bold
    return None


def paragraph_alignment(paragraph) -> str:
    value = paragraph.alignment
    if value is None and paragraph.style is not None:
        value = paragraph.style.paragraph_format.alignment
    return {WD_ALIGN_PARAGRAPH.LEFT: "左对齐", WD_ALIGN_PARAGRAPH.CENTER: "居中", WD_ALIGN_PARAGRAPH.RIGHT: "右对齐", WD_ALIGN_PARAGRAPH.JUSTIFY: "两端对齐"}.get(value, "未显式设置")


def line_spacing(paragraph):
    value = paragraph.paragraph_format.line_spacing
    if value is None and paragraph.style is not None:
        value = paragraph.style.paragraph_format.line_spacing
    if value is None:
        return None
    if isinstance(value, float) and value <= 3:
        return round(value, 2)
    try:
        return round(value.pt, 2)
    except AttributeError:
        return str(value)


def paragraph_numbering(paragraph) -> tuple[int, int] | None:
    candidates = []
    direct = paragraph._p.pPr.numPr if paragraph._p.pPr is not None else None
    if direct is not None:
        candidates.append(direct)
    style = paragraph.style
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        style_num = style.element.pPr.numPr if style.element.pPr is not None else None
        if style_num is not None:
            candidates.append(style_num)
        style = style.base_style
    for num_pr in candidates:
        num_id = int(num_pr.numId.val) if num_pr.numId is not None else -1
        ilvl = int(num_pr.ilvl.val) if num_pr.ilvl is not None else 0
        if num_id > 0 and ilvl >= 0:
            return num_id, ilvl
    return None


def numbering_templates(doc: Document) -> dict[tuple[int, int], str]:
    numbering_rel = next((rel for rel in doc.part.rels.values() if rel.reltype.endswith("/numbering")), None)
    if numbering_rel is None:
        return {}
    root = numbering_rel.target_part.element
    num_to_abstract = {}
    for num in root.findall(qn("w:num")):
        abstract_id = num.find(qn("w:abstractNumId"))
        if abstract_id is not None:
            num_to_abstract[int(num.get(qn("w:numId")))] = int(abstract_id.get(qn("w:val")))
    abstracts = {int(item.get(qn("w:abstractNumId"))): item for item in root.findall(qn("w:abstractNum"))}
    templates = {}
    for num_id, abstract_id in num_to_abstract.items():
        abstract = abstracts.get(abstract_id)
        if abstract is None:
            continue
        for level in abstract.findall(qn("w:lvl")):
            level_text = level.find(qn("w:lvlText"))
            if level_text is not None:
                templates[(num_id, int(level.get(qn("w:ilvl"))))] = clean(level_text.get(qn("w:val")))
    return templates


def automatic_numbering(doc: Document, paragraphs) -> Counter:
    templates = numbering_templates(doc)
    detected = Counter()
    for paragraph in paragraphs:
        numbering = paragraph_numbering(paragraph)
        if numbering is None:
            continue
        template = templates.get(numbering, "")
        if any(marker in template for marker in ("章", "节", "条")):
            detected[template] += 1
    return detected

def _paragraph_template(paragraph, templates) -> str | None:
    """取段落生效编号模板（含样式继承的 numPr）：(numId, ilvl) -> lvlText，如 '第%1章'。"""
    numbering = paragraph_numbering(paragraph)
    if numbering is None or templates is None:
        return None
    return templates.get(numbering, "")


def is_heading(paragraph, templates=None) -> bool:
    text = clean(paragraph.text)
    if paragraph.style.name.lower().startswith("heading"):
        return True
    if CHAPTER_RE.match(text) or SECTION_RE.match(text):
        return True
    template = _paragraph_template(paragraph, templates)
    return bool(template and any(marker in template for marker in ("章", "节")))


def is_chapter_heading(paragraph, templates=None) -> bool:
    """章标题判定：可见文本'第X章'，或 OOXML 自动编号模板含'章'且 ilvl=0（如 Heading 样式+numPr 渲染'第%1章'）。
    不满足时返回 False，交由调用方在报告缺失前完成 numbering.xml 模板核实，避免误报。"""
    text = clean(paragraph.text)
    if CHAPTER_RE.match(text):
        return True
    numbering = paragraph_numbering(paragraph)
    if numbering is None or templates is None:
        return False
    _, ilvl = numbering
    return ilvl == 0 and "章" in templates.get(numbering, "")


def compact(text: str) -> str:
    return re.sub(r"\s+", "", clean(text))


def title_core(text: str) -> str:
    return re.sub(r"[（(]试行[）)]$", "", compact(text)).strip()


def is_title_subject_line(text: str) -> bool:
    value = compact(text)
    return len(value) <= 60 and value.endswith(("股份有限公司", "集团有限公司", "有限责任公司", "分公司"))


def find_title(paragraphs) -> tuple[int, str]:
    for index, paragraph in enumerate(paragraphs[:20]):
        text = clean(paragraph.text)
        core = title_core(text)
        suffix = next((item for item in TITLE_SUFFIXES if core.endswith(item)), None)
        if not suffix or core.startswith("第") or len(core) > 80:
            continue
        start = index
        parts = [compact(text)]
        previous_index = index - 1
        while previous_index >= 0 and not clean(paragraphs[previous_index].text):
            previous_index -= 1
        if previous_index >= 0 and is_title_subject_line(paragraphs[previous_index].text):
            start = previous_index
            parts.insert(0, compact(paragraphs[previous_index].text))
        next_index = index + 1
        while next_index < len(paragraphs) and not clean(paragraphs[next_index].text):
            next_index += 1
        if next_index < len(paragraphs) and compact(paragraphs[next_index].text) in {"（试行）", "(试行)"}:
            parts.append(compact(paragraphs[next_index].text))
        return start, "".join(parts)
    return 0, clean(paragraphs[0].text) if paragraphs else "未识别"


def title_block(all_paragraphs, title_index: int, title: str):
    target = compact(title)
    selected = []
    combined = ""
    for index in range(title_index, min(title_index + 4, len(all_paragraphs))):
        paragraph = all_paragraphs[index]
        piece = compact(paragraph.text)
        if not piece:
            continue
        candidate = combined + piece
        if not target.startswith(candidate):
            break
        selected.append((index, paragraph))
        combined = candidate
        if combined == target:
            break
    return selected or [(title_index, all_paragraphs[title_index])]


def evidence(location: str, text: str, actual: str, required: str) -> dict:
    return {"location": location, "text": text[:220], "actual": actual, "required": required}


def naming_audit(path: Path, doc: Document, title_index: int, title: str) -> dict:
    structure = build_doc_structure(doc)
    formal_title = re.sub(r"[（(]试行[）)]", "", title).strip()
    suffix_hits = [suffix for suffix in TITLE_SUFFIXES if suffix in formal_title]
    declared = next((suffix for suffix in TITLE_SUFFIXES if formal_title.endswith(suffix)), None)
    deviations = []
    checks = []
    def add(rule_id, item, passed, actual, required, location="制度名称", severity="中", note=""):
        checks.append({"id": rule_id, "item": item, "passed": passed, "actual": actual, "required": required, "location": location, "severity": severity, "note": "" if passed else note})
        if not passed:
            deviations.append({"id": rule_id, "location": location, "actual": actual, "required": required, "severity": severity, "reason": note or item})
    add("N1", "层级词存在", bool(declared), title, "名称包含可识别层级词", severity="中", note="未识别制度层级词")
    add("N2", "层级词属于体系", declared in SUFFIXES, declared or "未识别", "章程/员工手册/管理办法/管理规定/实施细则/操作规程/指导手册/工作指引/议事规则/工作规则/工作细则", severity="中", note="层级词不在当前企业标准体系")
    has_subject_or_matter = len(re.sub(r"（试行）|\(试行\)", "", formal_title)) > (len(declared or "") + 1)
    if declared == "章程":
        has_subject_or_matter = len(formal_title) > len("章程")
    add("N3", "名称结构完整", has_subject_or_matter, formal_title, "适用主体（可省略）+管理事项+层级词", severity="中", note="未识别明确管理事项")
    add("N4", "层级词位置", bool(declared), title, "层级词位于名称末尾；试行标记可作为限定语", severity="中", note="层级词未位于名称末尾")
    repeated = len(suffix_hits) > 1
    add("N5", "无重复或复合层级词", not repeated, "、".join(suffix_hits), "原则上仅使用一个主要层级词", severity="中", note="存在重复或复合层级词")
    vague = any(word in formal_title for word in ["有关事项", "相关工作", "若干问题", "有关工作"])
    add("N6", "管理事项明确", not vague, formal_title, "事项边界清晰、避免通知式模糊表述", severity="中", note="名称包含边界不清表述")
    body_text = "\n".join(clean(p.text) for p in doc.paragraphs[title_index + 1:] if clean(p.text))
    matter = formal_title
    for suffix in TITLE_SUFFIXES:
        matter = matter.replace(suffix, "")
    matter = matter.replace("某集团有限公司", "").strip()
    matter_tokens = [token for token in re.split(r"[、，（）()\s]+", matter) if len(token) >= 2]
    token_hits = sum(1 for token in matter_tokens if token in body_text)
    ngram_hits = sum(1 for index in range(len(matter) - 1) if matter[index:index + 2] in body_text)
    charter_terms = ["公司", "股东大会", "董事会", "注册资本", "利润分配", "解散", "清算"]
    charter_consistency = declared == "章程" and sum(term in body_text for term in charter_terms) >= 4
    consistency = charter_consistency or (bool(matter_tokens) and (token_hits > 0 or ngram_hits >= 2))
    add("N7", "名称与正文对象一致", consistency, formal_title, "名称事项应在正文中得到实际规范", severity="高", note="名称事项与正文主题的可追溯性不足")
    management_match = declared in {"管理办法", "管理规定"}
    governance_match = declared in {"议事规则", "工作规则", "工作细则"} and any(word in body_text for word in ["委员", "会议", "表决", "回避", "会议记录"])
    charter_match = declared == "章程" and sum(term in body_text for term in charter_terms) >= 5
    handbook_terms = ["招聘", "试用期", "工作时间", "工资", "社会保险", "休假", "奖惩", "辞职"]
    handbook_match = declared == "员工手册" and sum(term in body_text for term in handbook_terms) >= 5
    if declared == "章程":
        level_required = "章程应系统规定公司基本事项、股东与治理机构、财务分配、重大变更、解散清算和修改生效"
    elif declared == "员工手册":
        level_required = "员工手册应系统覆盖入职、劳动合同、工时休假、薪酬社保、行为纪律、奖惩申诉和离职管理"
    else:
        level_required = "治理类规则应规定组成、职责、会议、表决、回避和记录"
    add("N8", "名称与制度层级匹配", management_match or governance_match or charter_match or handbook_match, declared or "未识别", level_required, severity="高", note="名称层级与正文内容定位不匹配")
    cover_text = "\n".join(clean(paragraph.text) for paragraph in doc.paragraphs[:title_index])
    organization_context = infer_organization_scope(formal_title, cover_text + "\n" + body_text)
    matched_orgs = organization_context["matched_organizations"]
    subject_actual = "、".join(item["canonical_name"] for item in matched_orgs) or "未识别特定组织名称"
    organization_scope = "company_wide" if declared == "章程" and "某集团有限公司" in formal_title else organization_context["scope"]
    named_companies = sorted(set(re.findall(r"[\u4e00-\u9fff]{2,30}(?:股份有限公司|有限责任公司|有限公司)", body_text)))
    named_companies = [name for name in named_companies if name not in formal_title]
    generic_employee_subject = declared == "员工手册" and formal_title.startswith("公司")
    employee_subject_issue = declared == "员工手册" and (generic_employee_subject or bool(named_companies))
    if employee_subject_issue:
        details = []
        if generic_employee_subject:
            details.append("正文标题使用‘公司’泛称")
        if named_companies:
            details.append("正文另出现=" + "、".join(named_companies[:4]))
        subject_actual = "；".join(details)
    add("N9", "适用主体表达", not employee_subject_issue, f"组织范围={organization_scope}；{subject_actual}", "特定主体使用现行组织名称；通用模板不得残留其他公司名称", severity="中", note="正式适用主体不明确或存在模板主体残留")
    work_markers = [marker for marker in ["-原版", "_通用版", "通用版", "模板", "草案", "征求意见稿"] if marker in path.stem]
    filename_flag = bool(work_markers)
    filename_title = path.stem
    for marker in work_markers:
        filename_title = filename_title.replace(marker, "")
    filename_title = filename_title.strip("_- ")
    filename_title = re.sub(r"[（(]试行[）)]", "", filename_title).strip()
    title_matches_filename = compact(filename_title) == compact(formal_title)
    carrier_pass = not filename_flag and title_matches_filename
    carrier_note = "文件名含工作版本标记：" + "、".join(work_markers) if filename_flag else "文件名与正文标题不一致" if not title_matches_filename else ""
    add("N10", "名称载体一致", carrier_pass, f"文件名={path.stem}；正文标题={formal_title}", "文件名不含“原版”“通用版”“模板”等工作标记，且与正文标题一致", severity="中", note=carrier_note)
    expected_self_name = {
        "管理办法": "本办法", "管理规定": "本规定", "实施细则": "本细则", "操作规程": "本规程",
        "章程": "本章程", "员工手册": "本手册", "指导手册": "本手册", "工作指引": "本指引", "议事规则": "本规则", "工作规则": "本规则", "工作细则": "本工作细则",
    }.get(declared)
    self_name_terms = ["本章程", "本办法", "本规定", "本细则", "本规程", "本手册", "本指引", "本规则", "本工作细则"]
    inconsistent_self_names = []
    if expected_self_name:
        for paragraph_index, paragraph in enumerate(doc.paragraphs[title_index + 1:], title_index + 2):
            paragraph_text = clean(paragraph.text)
            for term in self_name_terms:
                if term in paragraph_text and term != expected_self_name:
                    location = structure_label(structure[paragraph_index - 1]) or "正文"
                    inconsistent_self_names.append(f"{location}使用“{term}”")
    add("N11", "制度内部自称一致", not inconsistent_self_names, "；".join(inconsistent_self_names[:8]) or expected_self_name or "未识别", expected_self_name or "与正式制度名称一致", severity="中", note="正文中的制度自称与正式名称不一致")
    return {"title": title, "declared_level": declared or "其他", "checks": checks, "deviations": deviations}


def format_audit(path: Path, doc: Document, title_index: int, title: str) -> dict:
    all_paragraphs = list(doc.paragraphs)
    paragraphs = [paragraph for paragraph in all_paragraphs if clean(paragraph.text)]
    title_entries = title_block(all_paragraphs, title_index, title)
    title_paras = [paragraph for _, paragraph in title_entries]
    title_ids = {id(paragraph) for paragraph in title_paras}
    templates = numbering_templates(doc)
    numbering_meta = numbering_index(doc)

    def is_article_paragraph(paragraph) -> bool:
        if re.match(r"^第[一二三四五六七八九十百零〇0-9]+条", clean(paragraph.text)):
            return True
        numbering = paragraph_numbering(paragraph)
        return numbering is not None and "条" in templates.get(numbering, "")

    first_article_index = next((index for index, paragraph in enumerate(all_paragraphs) if is_article_paragraph(paragraph)), None)
    all_chapter_indexes = [index for index, paragraph in enumerate(all_paragraphs) if is_chapter_heading(paragraph, templates)]
    if first_article_index is not None:
        prior_chapters = [index for index in all_chapter_indexes if index < first_article_index]
        content_chapter_index = max(prior_chapters) if prior_chapters else title_index
    else:
        first_chapter_occurrences = [index for index in all_chapter_indexes if re.match(r"^第一章", clean(all_paragraphs[index].text))]
        content_chapter_index = first_chapter_occurrences[1] if len(first_chapter_occurrences) > 1 else all_chapter_indexes[0] if all_chapter_indexes else title_index
    chapter_paras = [
        paragraph for index, paragraph in enumerate(all_paragraphs)
        if index >= content_chapter_index and clean(paragraph.text) and is_chapter_heading(paragraph, templates)
    ]
    heading_ids = {
        id(paragraph) for index, paragraph in enumerate(all_paragraphs)
        if index >= content_chapter_index and clean(paragraph.text) and is_heading(paragraph, templates)
    }
    body_start_index = first_article_index if first_article_index is not None else content_chapter_index + 1
    body_paras = []
    attachment_started = False
    for index, paragraph in enumerate(all_paragraphs):
        text = clean(paragraph.text)
        if index < body_start_index or not text:
            continue
        if re.match(r"^附件\s*[一二三四五六七八九十百0-9]*\s*[：:]", text):
            attachment_started = True
        if attachment_started or id(paragraph) in heading_ids or id(paragraph) in title_ids or text in {"（试行）", "(试行)"}:
            continue
        body_paras.append(paragraph)

    def content_runs(paragraph):
        clause_match = re.match(r"^\s*第[一二三四五六七八九十百零〇0-9]+条\s*", paragraph.text)
        clause_end = clause_match.end() if clause_match else 0
        offset = 0
        for run in paragraph.runs:
            run_end = offset + len(run.text)
            if run.text.strip() and run_end > clause_end:
                yield run
            offset = run_end
    deviations = []
    checks = []
    def check(rule_id, item, passed, actual, required, location, severity="中", readable=True, note="", order=None):
        result = {"id": rule_id, "item": item, "passed": passed, "actual": actual, "required": required, "location": location, "severity": severity, "readable": readable, "note": note}
        checks.append(result)
        if readable and not passed:
            deviations.append({"id": rule_id, "location": location, "actual": actual, "required": required, "severity": severity, "reason": note or item, "order": order})
    title_fonts = sorted({font for paragraph in title_paras for run in paragraph.runs for font in run_fonts(run, paragraph)})
    title_sizes = sorted({size for paragraph in title_paras for run in paragraph.runs if (size := run_size(run, paragraph)) is not None})
    title_bold = all(run_bold(run, paragraph) is True for paragraph in title_paras for run in paragraph.runs if run.text.strip())
    title_alignments = [paragraph_alignment(paragraph) for paragraph in title_paras]
    title_location = "正文标题"
    title_order = title_index + 1
    title_font_pass = bool(title_fonts) and all(font in FONT_ALIASES["宋体"] for font in title_fonts)
    check("F1", "制度名称字体", title_font_pass, "、".join(title_fonts) or "未识别", "宋体", title_location, readable=bool(title_fonts), note="制度名称字体不是宋体", order=title_order)
    check("F2", "制度名称字号", title_sizes == [22.0], "、".join(map(str, title_sizes)) or "未识别", "22pt（二号）", title_location, readable=bool(title_sizes), note="制度名称字号不是二号", order=title_order)
    title_alignment_pass = bool(title_alignments) and all(value == "居中" for value in title_alignments)
    check("F3", "制度名称加粗和对齐", title_bold and title_alignment_pass, f"加粗={title_bold}；对齐={'/'.join(title_alignments)}", "加粗、居中", title_location, severity="低", note="制度名称未统一加粗或居中", order=title_order)
    heading_format_issues = []
    heading_layout_issues = []
    chapter_counters: dict[tuple[int, int], int] = {}
    for paragraph in chapter_paras:
        text = clean(paragraph.text)
        label = text
        numbering = paragraph_numbering(paragraph)
        if numbering is not None and not CHAPTER_RE.match(text):
            meta = numbering_meta.get(numbering)
            if meta is not None and "%1" in meta["template"]:
                chapter_counters[numbering] = chapter_counters.get(numbering, meta["start"] - 1) + 1
                label = f"{render_auto_label(meta, chapter_counters[numbering])} {text}".strip()
        runs = [run for run in paragraph.runs if run.text.strip()]
        fonts = sorted({font for run in runs for font in run_fonts(run, paragraph)})
        sizes = sorted({size for run in runs if (size := run_size(run, paragraph)) is not None})
        bold = all(run_bold(run, paragraph) is True for run in runs)
        alignment = paragraph_alignment(paragraph)
        if not fonts or not any(font in FONT_ALIASES["黑体"] for font in fonts) or sizes != [16.0]:
            heading_format_issues.append(f"{label}：字体={'/'.join(fonts) or '未识别'}，字号={sizes or '未识别'}")
        if not bold or alignment != FORMAT_STANDARD["chapter"]["alignment"]:
            heading_layout_issues.append(f"{label}：加粗={bold}，对齐={alignment}")
    chapter_first_order = all_paragraphs.index(chapter_paras[0]) + 1 if chapter_paras else None
    check("F4", "章标题字体字号", not heading_format_issues, "；".join(heading_format_issues[:5]) or "均为黑体16pt", "黑体、16pt", "章标题段落", readable=bool(chapter_paras), note="章标题字体或字号存在偏差", order=chapter_first_order)
    check("F5", "章标题加粗和对齐", not heading_layout_issues, "；".join(heading_layout_issues[:5]) or f"均已加粗并{FORMAT_STANDARD['chapter']['alignment']}", f"加粗、{FORMAT_STANDARD['chapter']['alignment']}", "章标题段落", severity="低", readable=bool(chapter_paras), note=f"章标题未统一加粗或{FORMAT_STANDARD['chapter']['alignment']}", order=chapter_first_order)
    body_fonts = Counter(font for paragraph in body_paras for run in content_runs(paragraph) for font in run_fonts(run, paragraph))
    body_sizes = Counter(size for paragraph in body_paras for run in content_runs(paragraph) if (size := run_size(run, paragraph)) is not None)
    body_font_pass = bool(body_fonts) and all(font in FONT_ALIASES["仿宋"] for font in body_fonts)
    body_size_pass = bool(body_sizes) and set(body_sizes) == {16.0}
    body_first_order = all_paragraphs.index(body_paras[0]) + 1 if body_paras else None
    check("F6", "正文字体", body_font_pass, "、".join(f"{k}({v})" for k, v in body_fonts.items()) or "未识别", "仿宋", "正文段落", severity="中", readable=bool(body_fonts), note="正文存在非仿宋字体", order=body_first_order)
    check("F7", "正文字号", body_size_pass, "、".join(f"{k}pt({v})" for k, v in body_sizes.items()) or "未识别", "16pt（三号）", "正文段落", severity="中", readable=bool(body_sizes), note="正文存在非三号字号", order=body_first_order)
    spacing_values = Counter(line_spacing(paragraph) for paragraph in body_paras)
    spacing_pass = all(value in {1.5, 28.0} for value in spacing_values if value is not None) and bool(spacing_values)
    check("F9", "正文行距", spacing_pass, f"行距={dict(spacing_values)}", "固定28磅或1.5倍", "正文段落", severity="低", readable=bool(body_paras), note="正文行距不符合基线", order=body_first_order)
    section = doc.sections[0] if doc.sections else None
    if section is None:
        check("F10", "页面尺寸和边距", False, "未识别", "A4；上/下2.54cm；左/右3.17cm", "页面设置", severity="低", readable=False)
    else:
        actual = f"纸张={section.page_width.inches:.2f}x{section.page_height.inches:.2f}in；边距=上{section.top_margin.cm:.2f}/下{section.bottom_margin.cm:.2f}/左{section.left_margin.cm:.2f}/右{section.right_margin.cm:.2f}cm"
        page_pass = abs(section.page_width.inches - 8.27) < .03 and abs(section.page_height.inches - 11.69) < .03 and abs(section.top_margin.cm - 2.54) < .05 and abs(section.bottom_margin.cm - 2.54) < .05 and abs(section.left_margin.cm - FORMAT_STANDARD["page"]["left_cm"]) < .05 and abs(section.right_margin.cm - FORMAT_STANDARD["page"]["right_cm"]) < .05
        check("F10", "页面尺寸和边距", page_pass, actual, "A4；上/下2.54cm；左/右3.17cm", "section:1", severity="低", note="页面尺寸或页边距不符合基线")
    header = " ".join(clean(p.text) for p in section.header.paragraphs) if section else ""
    footer = " ".join(clean(p.text) for p in section.footer.paragraphs) if section else ""
    header_xml = section.header._element.xml if section else ""
    footer_xml = section.footer._element.xml if section else ""
    page_field_location = "页眉" if "PAGE" in header_xml else "页脚" if "PAGE" in footer_xml else "未发现"
    check("F11", "页眉页脚和页码", page_field_location != "未发现", f"页眉文本={header or '空'}；页脚文本={footer or '空'}；页码域={page_field_location}", "页眉或页脚设置统一页码", "页眉/页脚", severity="低", note="未发现页码域")
    numbering = [clean(paragraph.text).split()[0] for paragraph in paragraphs if re.match(r"^第[一二三四五六七八九十百零〇0-9]+(?:章|节|条)", clean(paragraph.text))]
    auto_numbering = automatic_numbering(doc, paragraphs)
    numbering_actual = []
    if numbering:
        numbering_actual.append("可见编号=" + "、".join(numbering[:12]))
    if auto_numbering:
        numbering_actual.append("自动编号模板=" + "、".join(f"{template}({count}处)" for template, count in auto_numbering.items()))
    check("F12", "编号层级可识别", bool(numbering or auto_numbering), "；".join(numbering_actual) or "未识别章/节/条编号", "可见编号或OOXML自动编号均可识别", "正文编号", severity="低", readable=True, note="未识别规范的章/节/条编号")
    # 偏差行按文档顺序排列：无文档位置的（页面/页眉等）置后保持规则顺序
    deviations.sort(key=lambda d: (d.get("order") is None, d.get("order") if isinstance(d.get("order"), int) else 0))
    return {"page": {"section_count": len(doc.sections), "header": header, "footer": footer}, "statistics": {"title_fonts": title_fonts, "title_sizes": title_sizes, "body_fonts": dict(body_fonts), "body_sizes": dict(body_sizes), "body_line_spacing": dict(spacing_values), "chapter_count": len(chapter_paras), "body_paragraph_count": len(body_paras)}, "checks": checks, "deviations": deviations}


def report(audits: list[dict]) -> str:
    lines = ["# 制度规范性审核报告", "", "> 审核基线：命名规范体系和公司制度格式要求。", "> 审核范围：制度名称、层级匹配、字体字号、段落、页面设置、页眉页脚和编号。", ""]
    for audit in audits:
        naming = audit["naming"]; formatting = audit["format"]
        lines += [f"## {audit['document']['file_name']}", "", f"- 制度名称：{naming['title']}", f"- 名称层级：{naming['declared_level']}", f"- 命名偏差：{len(naming['deviations'])}项", f"- 格式偏差：{len(formatting['deviations'])}项", ""]
        lines += ["### 命名审核", "", "| 编号 | 检查项 | 状态 | 实际 | 要求 | 严重程度 | 说明 |", "|---|---|---|---|---|---|---|"]
        for item in naming["checks"]:
            lines.append(f"| {item['id']} | {item['item']} | {'通过' if item['passed'] else '不通过'} | {item['actual']} | {item['required']} | {item['severity']} | {item['note']} |")
        lines += ["", "### 格式审核偏差", "", "| 编号 | 位置 | 检查项 | 实际 | 要求 | 严重程度 | 建议 |", "|---|---|---|---|---|---|---|"]
        if formatting["deviations"]:
            for item in formatting["deviations"]:
                lines.append(f"| {item['id']} | {item['location']} | {item['reason']} | {item['actual']} | {item['required']} | {item['severity']} | 建议按要求调整该位置格式。 |")
        else:
            lines.append("| - | - | 未发现偏差 | - | - | - | - |")
        lines += ["", "### 格式统计", "", "```json", json.dumps(formatting["statistics"], ensure_ascii=False, indent=2), "```", ""]
    return "\n".join(lines)


def run(paths: list[Path], output_dir: Path) -> list[dict]:
    output_dir.mkdir(parents=True, exist_ok=True)
    audits = []
    for path in paths:
        doc = Document(str(path))
        title_index, title = find_title(doc.paragraphs)
        item = {"audit_info": {"parsed_at": datetime.now().isoformat(timespec="seconds"), "basis": ["references/company_drafting_standard.md", "scripts/audit_normative.py/FORMAT_STANDARD", "references/organization_structure.json"], "profile": QUICK_AUDIT_PROFILE}, "document": {"file_name": path.name, "source_file": str(path.resolve())}, "naming": naming_audit(path, doc, title_index, title), "format": format_audit(path, doc, title_index, title), "parse_issues": []}
        audits.append(item)
        safe = re.sub(r"[\\/:*?\"<>|]", "_", path.stem)
        (output_dir / f"{safe}.json").write_text(json.dumps(item, ensure_ascii=False, indent=2), encoding="utf-8")
    report_name = "制度规范性审核报告.md"
    summary_name = "制度规范性审核汇总.json"
    (output_dir / report_name).write_text(report(audits), encoding="utf-8")
    (output_dir / summary_name).write_text(json.dumps({"audits": audits}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"documents": len(audits), "output_dir": str(output_dir), "naming_deviations": [len(x["naming"]["deviations"]) for x in audits], "format_deviations": [len(x["format"]["deviations"]) for x in audits]}, ensure_ascii=False, indent=2))
    return audits


def main() -> None:
    parser = argparse.ArgumentParser(description="按企业命名和格式基线审核 DOCX 规范性")
    parser.add_argument("documents", nargs="+", type=Path)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    args = parser.parse_args()
    run(args.documents, args.output_dir)


if __name__ == "__main__":
    main()



