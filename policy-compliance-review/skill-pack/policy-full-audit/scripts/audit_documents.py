from __future__ import annotations

import argparse
import json
import re
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from audit_common import (
    QUICK_AUDIT_PROFILE, build_doc_structure, concern_sort_key,
    infer_organization_scope, numbering_index, paragraph_numbering,
    structure_label,
)

SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_ROOT = SCRIPT_DIR.parent
DEFAULT_OUTPUT_DIR = Path.cwd() / "审核输出" / "completeness"
LEVEL_SUFFIXES = ["章程", "员工手册", "管理办法", "管理规定", "实施细则", "操作规程", "指导手册", "工作指引", "议事规则", "工作规则", "工作细则"]

@dataclass
class ParagraphRecord:
    id: str
    text: str
    style: str
    heading_kind: str | None
    bold: bool | None
    font_names: list[str]
    font_sizes_pt: list[float]
    alignment: str
    numbering_template: str | None = None  # OOXML 自动编号 lvlText 模板，如 "第%1章"/"第%1条"；无自动编号为 None
    location_label: str = ""  # 章条坐标，如 "第一章第三条"/"第一章"；标题区等无章条归属为空串

@dataclass
class TableRecord:
    id: str
    text: str

# 通用必备条款规则（与 references/mandatory_clause_catalog.md 二、通用管理制度必备条款 12 条对齐，2026-09-04 升级）
# 字段说明：
#   groups          内容要素分组（关键词为检索召回词，不直接等同于通过；多组用于区分"实质内容"与"仅提及"）
#   match_mode      any=任一组命中即存在；all=全部组命中才存在，缺组为部分；priority=groups[0] 命中即存在，仅后组命中为部分
#   applies_when    catalog 适用条件（人类可读）
#   applies_triggers 适用触发器：全文命中任一即启用该审核点；为 None 表示无条件适用（全部正式制度/管理性制度原则上适用）
GENERIC_RULES = [
    {"id": "M-GEN-001", "item": "制定目的和依据", "severity": "中", "applies_when": "全部正式制度", "applies_triggers": None, "match_mode": "any", "groups": [["为规范", "为加强", "为明确", "为完善", "为有效", "为了", "特制定", "制定本", "根据", "依据", "目的"]]},
    {"id": "M-GEN-002", "item": "适用范围", "severity": "高", "applies_when": "全部正式制度", "applies_triggers": None, "match_mode": "any", "groups": [["适用范围", "适用对象", "适用单位", "适用于", "用于指导", "本办法用于", "本规定用于"]], "strong": ["适用范围", "适用对象", "适用单位", "适用于"]},
    {"id": "M-GEN-003", "item": "术语定义", "severity": "中", "applies_when": "存在专业术语或自定义概念", "applies_triggers": ["以下简称", "简称", "术语", "定义", "是指"], "match_mode": "priority", "groups": [["术语", "定义", "是指", "即指", "系指", "本办法所称", "本规定所称"], ["以下简称", "简称"]]},
    {"id": "M-GEN-004", "item": "职责分工", "severity": "高", "applies_when": "全部正式制度", "applies_triggers": None, "match_mode": "any", "groups": [["负责", "归口", "职责", "分工", "权限"]]},
    {"id": "M-GEN-005", "item": "管理条件和标准", "severity": "高", "applies_when": "涉及准入、评审、验收、审批或考核", "applies_triggers": ["准入", "评审", "验收", "审批", "考核"], "match_mode": "priority", "groups": [["应符合", "应满足", "不得低于", "不得高于", "必须达到", "标准如下", "条件如下", "准入条件", "详见附件", "按照附件"], ["标准", "条件"]]},
    {"id": "M-GEN-006", "item": "管理程序", "severity": "高", "applies_when": "涉及业务办理或审批", "applies_triggers": ["申请", "受理", "审核", "审批", "流程", "办理"], "match_mode": "all", "groups": [["申请", "受理", "提出"], ["审核", "审批", "批准", "审定"], ["反馈", "备案", "跟踪"]]},
    {"id": "M-GEN-007", "item": "权限体系", "severity": "高", "applies_when": "涉及决策、授权或资金资源", "applies_triggers": ["决策", "授权", "额度", "资金", "预算", "投资"], "match_mode": "any", "groups": [["权限", "授权", "额度", "超权限", "转授权", "例外审批"]]},
    {"id": "M-GEN-008", "item": "时限要求", "severity": "中", "applies_when": "时间影响权利义务或流程效率", "applies_triggers": ["时限", "期限", "工作日", "日内", "日前", "小时内", "天内"], "match_mode": "any", "groups": [["时限", "期限", "工作日", "日内", "日前", "小时", "天内", "天前"]]},
    {"id": "M-GEN-009", "item": "记录和归档", "severity": "中", "applies_when": "需要留痕、追溯或接受审计", "applies_triggers": ["留痕", "追溯", "审计", "台账", "归档", "档案"], "match_mode": "any", "groups": [["记录", "归档", "台账", "留痕", "档案", "保存"]]},
    {"id": "M-GEN-010", "item": "监督检查", "severity": "中", "applies_when": "管理性制度原则上适用", "applies_triggers": None, "match_mode": "any", "groups": [["监督检查", "监督管理", "检查", "整改"]]},
    {"id": "M-GEN-011", "item": "责任处理", "severity": "中", "applies_when": "设置义务或禁止事项", "applies_triggers": ["禁止", "不得", "义务", "罚款", "处罚", "处分", "吊销", "没收", "赔偿", "违规", "问责", "追究", "移交"], "match_mode": "all", "groups": [["违规", "责任认定", "责任追究", "问责", "移交", "移送"], ["罚款", "处罚", "处分", "吊销", "没收", "赔偿", "违约金", "保证金"]]},
    {"id": "M-GEN-012", "item": "解释和施行", "severity": "中", "applies_when": "全部正式制度", "applies_triggers": None, "match_mode": "any", "groups": [["负责解释", "解释权", "解释", "施行", "生效", "废止", "修订"]]},
]

MANAGEMENT_CHAPTERS = [
    ("S-MA-001", "总则", ["总则", "制定目的", "制定依据", "基本原则"], "必须", "中"),
    ("S-MA-002", "适用范围", ["适用范围", "适用对象", "适用单位", "适用于", "用于指导", "本办法用于", "本规定适用于"], "必须", "中"),
    ("S-MA-003", "术语定义", ["术语和定义", "术语定义", "本办法所称", "本规定所称"], "条件必须", "中"),
    ("S-MA-004", "职责分工", ["职责分工", "管理职责", "工作职责", "权限"], "必须", "高"),
    ("S-MA-005", "具体管理规范", ["管理要求", "管理规范", "具体管理", "合作管理", "质量管理", "监督管理"], "必须", "高"),
    ("S-MA-006", "监督检查", ["监督检查", "监督管理", "检查", "整改"], "应当", "中"),
    ("S-MA-007", "考核或责任", ["考核", "责任追究", "责任认定", "违规处理", "问责"], "条件必须", "中"),
    ("S-MA-008", "附则", ["附则", "解释", "施行", "生效", "执行", "废止", "修订"], "必须", "中"),
]

PROCUREMENT_AGENT_RULES = [
    ("M-PA-001", "采购代理机构职责和权限", ["采购代理机构", "采购部门", "职责", "权限"], "B", "高"),
    ("M-PA-002", "采购代理机构选聘和委托", ["选择", "寻源", "引入", "委托服务合同", "委托代理协议"], "B", "高"),
    ("M-PA-003", "代理服务合同和履约", ["委托服务合同", "合同", "代理服务", "履约"], "B", "高"),
    ("M-PA-004", "代理服务收费", ["服务费", "收费", "费率", "费用标准"], "B", "中"),
    ("M-PA-005", "采购代理业务保密和廉洁", ["保密", "廉洁", "利益冲突", "回避", "不得"], "B", "高"),
    ("M-PA-006", "采购代理机构考核评价", ["考核", "评价", "阶段性评估", "后评估"], "B", "中"),
    ("M-PA-007", "异议投诉和监督处理", ["异议", "投诉", "举报", "监督检查", "整改"], "B", "高"),
    ("M-PA-008", "采购代理机构档案和记录", ["档案", "记录", "归档", "留痕"], "B", "中"),
]

MEETING_RULES = [
    ("M-MTG-001", "组织性质和依据", [["审计委员会", "专门工作机构"], ["根据", "公司章程"], ["独立开展", "对董事会负责", "工作原则"]], "B", "高"),
    ("M-MTG-002", "组成和任期", [["三名以上", "独立非执行董事"], ["主任委员"], ["任期"], ["工作小组", "日常办事机构"], ["更换", "补足委员"]], "B", "高"),
    ("M-MTG-003", "职责和议事范围", [["职责权限"], ["审阅", "监察", "检讨", "监督"], ["董事会", "汇报", "建议"]], "B", "高"),
    ("M-MTG-004", "会议类型和频次", [["定期会议"], ["临时会议"], ["每年至少", "每半年至少"], ["提议", "召开条件"]], "B", "中"),
    ("M-MTG-005", "召集和通知", [["主任委员", "召集"], ["通知"], ["会议材料", "会议议程"], ["五天", "三天", "一天"]], "B", "中"),
    ("M-MTG-006", "出席人数和委托", [["过半数出席"], ["委托"], ["列席"], ["不能出席", "未出席"]], "B", "高"),
    ("M-MTG-007", "利益冲突和回避", [["回避"], ["关联方", "利益冲突"], ["利益申报", "关联关系申报", "披露利益"], ["不计入", "人数计算", "表决权"]], "B", "高"),
    ("M-MTG-008", "审议和表决", [["表决"], ["半数以上", "过半数"], ["现场会议", "书面议案"], ["弃权"], ["异议", "不同意见"]], "B", "高"),
    ("M-MTG-009", "会议记录和纪要", [["会议记录"], ["签字"], ["保密"], ["保存", "归档"], ["初稿", "定稿"]], "B", "高"),
    ("M-MTG-010", "决议执行和督办", [["报送公司董事会", "决议执行", "执行主体"], ["督办", "跟踪"], ["反馈"], ["决议调整", "执行调整", "纠偏"]], "B", "中"),
]

MEETING_CHAPTER_RULES = [
    ("S-MTG-001", "组织定位", [["总则"], ["审计委员会", "专门工作机构"], ["根据", "公司章程"]], "必须", "中"),
    ("S-MTG-002", "组成及任期", [["人员组成"], ["主任委员"], ["任期"], ["更换", "补足委员"]], "必须", "高"),
    ("S-MTG-003", "职责权限", [["职责权限"], ["审阅", "监察", "监督"], ["董事会", "汇报"]], "必须", "高"),
    ("S-MTG-004", "会议召集", [["议事规则"], ["定期会议"], ["临时会议"], ["通知"]], "必须", "中"),
    ("S-MTG-005", "议题和材料", [["议题"], ["会议材料"], ["会前审查", "预审", "审核材料"]], "应当", "中"),
    ("S-MTG-006", "出席和回避", [["出席"], ["委托"], ["列席"], ["回避"]], "必须", "高"),
    ("S-MTG-007", "审议和表决", [["表决"], ["半数以上", "过半数"], ["异议", "不同意见"]], "必须", "高"),
    ("S-MTG-008", "记录和执行", [["会议记录"], ["保存", "归档"], ["报送公司董事会"], ["督办", "跟踪", "决议执行"]], "必须", "中"),
    ("S-MTG-009", "附则", [["附则"], ["解释"], ["生效"], ["修订"]], "必须", "中"),
]

COMPANY_CHARTER_RULES = [
    ("M-CHT-001", "公司基本事项", [["公司名称", "公司的名称"], ["住所"], ["经营宗旨", "经营范围"]], "章程核心", "高"),
    ("M-CHT-002", "注册资本和股份结构", [["注册资本"], ["股份总数", "总股本", "股本结构"], ["股票", "股份"]], "章程核心", "高"),
    ("M-CHT-003", "股东权利义务和股东名册", [["股东的权利"], ["股东的义务", "股东义务", "股东承担", "承担下列义务"], ["股东名册"]], "章程核心", "高"),
    ("M-CHT-004", "股东会及其议事机制", [["股东大会", "股东会"], ["召集"], ["提案", "通知"], ["表决", "决议"]], "章程核心", "高"),
    ("M-CHT-005", "董事会治理安排", [["董事会"], ["董事会职权", "董事会行使", "董事会对股东大会负责", "行使下列主要职权"], ["董事任期", "任期"], ["董事会会议", "表决"]], "章程核心", "高"),
    ("M-CHT-006", "经理层设置和职权", [["总经理", "经理层", "经营管理机构"], ["职权", "行使下列职权"], ["董事会", "聘任", "解聘"]], "章程核心", "高"),
    ("M-CHT-007", "监督机构治理安排", [["监事会", "审计委员会"], ["监督", "监察"], ["组成", "任期"]], "章程核心", "高"),
    ("M-CHT-008", "董监高资格和义务", [["董事、监事和高级管理人员", "董事、监事、高级管理人员"], ["资格", "任职"], ["义务", "忠实", "勤勉"]], "章程核心", "高"),
    ("M-CHT-009", "财务会计、审计和利润分配", [["财务会计"], ["利润分配", "股利"], ["会计师事务所", "审计"]], "章程核心", "高"),
    ("M-CHT-010", "通知和公告", [["通知"], ["公告"]], "章程核心", "中"),
    ("M-CHT-011", "合并分立及资本变动", [["合并"], ["分立"], ["增资", "减资", "增加注册资本", "减少注册资本"]], "章程核心", "高"),
    ("M-CHT-012", "解散和清算", [["解散"], ["清算"]], "章程核心", "高"),
    ("M-CHT-013", "章程修改程序", [["本章程的修改", "修改章程", "章程修改"], ["股东大会", "股东会"], ["特别决议", "决议"]], "章程核心", "高"),
    ("M-CHT-014", "解释、生效和附则", [["解释"], ["生效"], ["附则"]], "章程核心", "中"),
]

COMPANY_CHARTER_CHAPTERS = [
    ("S-CHT-001", "总则", ["总则"], "必须", "中"),
    ("S-CHT-002", "经营宗旨和范围", ["经营宗旨和范围", "经营范围"], "必须", "高"),
    ("S-CHT-003", "股份和注册资本", ["股份和注册资本", "注册资本", "股本结构"], "必须", "高"),
    ("S-CHT-004", "股东权利义务", ["股东的权利和义务", "股东权利义务"], "必须", "高"),
    ("S-CHT-005", "股东会治理", ["股东大会", "股东会"], "必须", "高"),
    ("S-CHT-006", "董事会治理", ["董事会"], "必须", "高"),
    ("S-CHT-007", "经理层或经营管理机构", ["经营管理机构", "经理层", "总经理"], "必须", "高"),
    ("S-CHT-008", "监督机构", ["监事会", "审计委员会"], "必须", "高"),
    ("S-CHT-009", "董监高资格和义务", ["董事、监事和高级管理人员的资格和义务", "高级管理人员的资格和义务"], "必须", "高"),
    ("S-CHT-010", "财务会计、审计和利润分配", ["财务会计制度与利润分配", "会计师事务所", "审计"], "必须", "高"),
    ("S-CHT-011", "合并分立、解散清算和章程修改", ["公司的合并", "公司解散", "本章程的修改程序", "章程修改"], "必须", "高"),
    ("S-CHT-012", "附则", ["附则"], "必须", "中"),
]

def clean(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()

def heading_kind(text: str, style: str, template: str | None = None) -> str | None:
    if style.lower().startswith("heading"):
        return style
    if re.match(r"^第[一二三四五六七八九十百]+章", text):
        return "chapter"
    if re.match(r"^第[一二三四五六七八九十百]+节", text):
        return "section"
    if re.match(r"^第[一二三四五六七八九十百]+条", text):
        return "article"
    if re.match(r"^(附件|附录)[A-Z一二三四五六七八九十]", text):
        return "appendix"
    if template:
        # OOXML 自动编号模板兜底：模板含"章/节/条"即视为该级标题，禁止当作普通段落或误报缺失
        if "章" in template:
            return "chapter"
        if "节" in template:
            return "section"
        if "条" in template:
            return "article"
    return None

def alignment_name(value) -> str:
    return {WD_ALIGN_PARAGRAPH.LEFT: "left", WD_ALIGN_PARAGRAPH.CENTER: "center", WD_ALIGN_PARAGRAPH.RIGHT: "right", WD_ALIGN_PARAGRAPH.JUSTIFY: "justify"}.get(value, "unspecified")

EMPLOYEE_HANDBOOK_RULES = [
    ("M-LAB-001", "劳动合同期限", [["劳动合同"], ["合同期限", "固定期限", "无固定期限"]], "A/B", "高"),
    ("M-LAB-002", "工作内容和地点", [["工作内容", "岗位职责", "本职工作"], ["工作地点", "工作场所"], ["调职", "岗位调整", "工作调动"]], "A/B", "高"),
    ("M-LAB-003", "工作时间和休息休假", [["工作时间"], ["休息", "公休日"], ["法定假日", "法定节假日"], ["年休假", "病假", "婚假", "产假"]], "A/B", "高"),
    ("M-LAB-004", "劳动报酬", [["工资", "报酬"], ["发工资日", "工资支付", "发放"], ["基本工资", "工资待遇"], ["扣除", "扣款"]], "A/B", "高"),
    ("M-LAB-005", "社会保险", [["社会保险", "社保"], ["缴纳", "参保"], ["养老保险", "医疗保险", "工伤保险"]], "A", "高"),
    ("M-LAB-006", "劳动保护和职业健康", [["安全生产", "安全卫生", "安全和健康"], ["劳动保护", "职业健康", "职业病", "职业危害"], ["事故", "灾害", "应急"]], "A/B", "高"),
    ("M-LAB-007", "试用期", [["试用期"], ["一个月", "二个月", "六个月"], ["转正", "试用考核"], ["不符合录用条件"]], "A", "高"),
    ("M-LAB-008", "规章制度民主程序", [["职工代表大会", "全体职工讨论", "民主程序"], ["工会", "职工代表"], ["协商确定", "平等协商"]], "A", "高"),
    ("M-LAB-009", "公示告知和签收", [["公示", "告知", "通知"], ["签收", "员工签名"], ["生效", "颁布"]], "A/B", "高"),
    ("M-LAB-010", "解除终止和离职", [["解除劳动合同", "终止聘用关系", "辞退"], ["辞职"], ["提前一个月", "提前30日", "提前3日"], ["经济补偿"], ["离职证明"]], "A/B", "高"),
    ("M-LAB-011", "培训服务期和竞业限制", [["服务期", "专项培训费用"], ["竞业限制", "竞业禁止"], ["违约金", "补偿金"]], "A", "中"),
    ("M-LAB-012", "考勤和加班管理", [["考勤"], ["加班"], ["批准", "审批"], ["加班费", "调休", "补休", "加班记录"]], "B", "中"),
    ("M-LAB-013", "员工奖惩和申诉", [["奖励", "嘉奖", "表彰"], ["惩罚", "处分", "记过", "降级"], ["申诉", "申辩"], ["批准", "审批"], ["通知", "送达"]], "A/B", "高"),
    ("M-LAB-014", "平等就业和反歧视", [["平等就业", "公平就业", "不得歧视"], ["性别", "民族", "宗教", "残疾"]], "A/B", "高"),
    ("M-LAB-015", "员工个人信息", [["个人信息", "员工档案", "个人资料"], ["收集", "缴验", "提供"], ["保存期限", "查阅权限", "信息安全", "隐私"]], "A", "高"),
    ("M-LAB-016", "劳动争议处理", [["劳动争议"], ["调解", "仲裁", "诉讼"], ["申诉渠道", "争议处理"]], "B", "中"),
]

EMPLOYEE_HANDBOOK_CHAPTERS = [
    ("S-HB-001", "前言、依据和适用范围", ["前言", "总则", "适用范围"], "必须", "中"),
    ("S-HB-002", "员工守则和行为规范", ["员工守则", "行为准则", "工作纪律"], "必须", "高"),
    ("S-HB-003", "招聘、入职和试用", ["招聘", "入职", "试用"], "必须", "高"),
    ("S-HB-004", "工作时间和考勤", ["工作", "工作时间", "考勤"], "必须", "高"),
    ("S-HB-005", "工资福利和社会保险", ["工资与福利", "工资", "社会保险"], "必须", "高"),
    ("S-HB-006", "请假和休假", ["请假", "休假", "年休假"], "必须", "高"),
    ("S-HB-007", "加班管理", ["加班"], "必须", "中"),
    ("S-HB-008", "培训、调职和发展", ["培训", "调职", "岗位调整"], "应当", "中"),
    ("S-HB-009", "保密、安全和职业健康", ["保密", "安全质量", "安全生产", "职业健康"], "必须", "高"),
    ("S-HB-010", "考核、奖励和纪律处分", ["考核", "奖惩", "处分"], "必须", "高"),
    ("S-HB-011", "辞退、辞职和离职", ["辞退", "辞职", "解除劳动合同"], "必须", "高"),
    ("S-HB-012", "申诉和劳动争议", ["申诉", "劳动争议", "争议处理"], "必须", "高"),
    ("S-HB-013", "解释、生效和签收", ["其他", "解释", "生效", "签收"], "必须", "中"),
]

def title_core(text: str) -> str:
    return re.sub(r"[（(]试行[）)]$", "", re.sub(r"\s+", "", text or "")).strip()


def is_title_subject_line(text: str) -> bool:
    value = re.sub(r"\s+", "", text or "")
    return len(value) <= 60 and value.endswith(("股份有限公司", "集团有限公司", "有限责任公司", "分公司"))


def detect_body_title(paragraphs: list[ParagraphRecord]) -> str:
    for index, paragraph in enumerate(paragraphs[:20]):
        core = title_core(paragraph.text)
        suffix = next((item for item in LEVEL_SUFFIXES if core.endswith(item)), None)
        if not suffix or core.startswith("第") or len(core) > 80:
            continue
        start = index
        parts = [re.sub(r"\s+", "", paragraph.text)]
        if index > 0 and is_title_subject_line(paragraphs[index - 1].text):
            start = index - 1
            parts.insert(0, re.sub(r"\s+", "", paragraphs[start].text))
        if index + 1 < len(paragraphs) and re.sub(r"\s+", "", paragraphs[index + 1].text) in {"（试行）", "(试行)"}:
            parts.append(re.sub(r"\s+", "", paragraphs[index + 1].text))
        return "".join(parts)
    return paragraphs[0].text if paragraphs else "未识别"

def _run_fonts(run, paragraph) -> list[str]:
    """run 实际使用字体：显式 rFonts（eastAsia/ascii/hAnsi）优先，无显式设置时回退段落样式字体。
    注意：仅条号（如"第七条"）显式黑体、正文继承样式的段落，会同时列出黑体与样式字体，
    不得将 font_names 理解为"整段同一种字体"。"""
    rpr = run._r.rPr
    if rpr is not None and rpr.rFonts is not None:
        rfonts = rpr.rFonts
        direct = [v for v in [rfonts.get(qn("w:eastAsia")), rfonts.get(qn("w:ascii")), rfonts.get(qn("w:hAnsi"))] if v]
        if direct:
            return sorted(set(direct))
    if run.font.name:
        return [run.font.name]
    style = paragraph.style
    return [style.font.name] if style is not None and style.font.name else []


def extract_docx(path: Path) -> dict:
    doc = Document(str(path))
    meta_index = numbering_index(doc)
    structure = build_doc_structure(doc)
    paragraphs = []
    for index, paragraph in enumerate(doc.paragraphs, 1):
        text = clean(paragraph.text)
        if not text:
            continue
        numbering = paragraph_numbering(paragraph)
        meta = meta_index.get(numbering) if numbering is not None else None
        template = meta["template"] if meta else None
        font_names = sorted({font for run in paragraph.runs if run.text.strip() for font in _run_fonts(run, paragraph)})
        font_sizes = sorted({round(run.font.size.pt, 2) for run in paragraph.runs if run.font.size is not None})
        bold_values = [run.bold for run in paragraph.runs if run.text.strip()]
        bold = None if not bold_values else all(value is True for value in bold_values)
        paragraphs.append(ParagraphRecord(f"para:{index}", text, paragraph.style.name, heading_kind(text, paragraph.style.name, template), bold, font_names, font_sizes, alignment_name(paragraph.alignment), template, structure_label(structure[index - 1])))
    tables = []
    for table_index, table in enumerate(doc.tables, 1):
        rows = [" | ".join(clean(cell.text) for cell in row.cells) for row in table.rows]
        tables.append(TableRecord(f"table:{table_index}", "\n".join(rows)))
    all_text = "\n".join(item.text for item in paragraphs) + "\n" + "\n".join(item.text for item in tables)
    title = detect_body_title(paragraphs)
    return {"path": str(path.resolve()), "file_name": path.name, "title": title, "paragraphs": paragraphs, "tables": tables, "all_text": all_text, "parse_issues": []}

def serialize_doc(doc: dict) -> dict:
    result = dict(doc)
    result["paragraphs"] = [asdict(item) for item in doc["paragraphs"]]
    result["tables"] = [asdict(item) for item in doc["tables"]]
    result.pop("all_text", None)
    return result

def classify(doc: dict) -> dict:
    name = doc["title"]
    if "员工手册" in name:
        primary = "employee_handbook"
    elif "章程" in name:
        primary = "company_charter"
    elif any(word in name for word in ["董事会", "委员会", "议事规则", "工作规则", "工作细则"]):
        primary = "meeting_governance"
    elif "采购代理机构" in name:
        primary = "procurement_agent"
    elif any(x in name for x in ["对外合作", "合作管理", "重点实验室", "科研合作"]):
        primary = "cooperation"
    elif any(x in name for x in ["信息化专业", "信息化质量", "质量管理办法"]):
        primary = "information_quality"
    else:
        primary = "generic"
    level = next((suffix for suffix in LEVEL_SUFFIXES if suffix in doc["title"]), "其他")
    organization_context = infer_organization_scope(doc["title"], doc["all_text"])
    organization_scope = "company_wide" if primary == "company_charter" else organization_context["scope"]
    return {
        "primary_type": primary,
        "type_confidence": "high" if primary != "generic" else "low",
        "declared_level": level,
        "chapter_profile": "handbook" if primary == "employee_handbook" else "charter" if primary == "company_charter" else "meeting" if primary == "meeting_governance" else "management",
        "custom_type": primary not in {"procurement_agent", "meeting_governance", "company_charter", "employee_handbook"},
        "organization_scope": organization_scope,
        "matched_organizations": organization_context["matched_organizations"],
    }

def evidence_for(doc: dict, keywords: list[str], limit: int = 5) -> list[dict]:
    results = []
    for item in doc["paragraphs"]:
        compact_text = re.sub(r"\s+", "", item.text)
        hits = [key for key in keywords if re.sub(r"\s+", "", key) in compact_text]
        if hits:
            results.append({"location": _paragraph_location(doc, item), "text": item.text[:220], "matched": hits[:6], "order": int(item.id.split(":")[1])})
            if len(results) >= limit:
                return results
    for item in doc["tables"]:
        compact_text = re.sub(r"\s+", "", item.text)
        hits = [key for key in keywords if re.sub(r"\s+", "", key) in compact_text]
        if hits:
            table_index = int(item.id.split(":")[1])
            results.append({"location": f"附表{table_index}", "text": item.text[:220], "matched": hits[:6], "order": 900000 + table_index})
            if len(results) >= limit:
                return results
    return results


def _paragraph_location(doc: dict, item) -> str:
    """位置显示统一使用章条坐标（如 '第一章第三条'）；标题区等无章条归属的给语义位置，不显示段落号。"""
    if item.location_label:
        return item.location_label
    title = doc.get("title") or ""
    if title and re.sub(r"\s+", "", title) in re.sub(r"\s+", "", item.text):
        return "正文标题"
    return "正文"

def grouped_status(doc: dict, groups: list[list[str]]) -> tuple[str, list[dict], list[list[str]]]:
    evidence = []
    missing_groups = []
    for group in groups:
        group_evidence = evidence_for(doc, group)
        if group_evidence:
            for item in group_evidence:
                if item not in evidence:
                    evidence.append(item)
        else:
            missing_groups.append(group)
    if not evidence:
        return "缺失", [], missing_groups
    if missing_groups:
        return "部分", evidence[:12], missing_groups
    return "存在", evidence[:12], []


def check_grouped_rules(doc: dict, rules: list[tuple], chapter: bool = False) -> list[dict]:
    results = []
    for rule_id, item, groups, source_or_level, severity in rules:
        status, evidence, missing_groups = grouped_status(doc, groups)
        missing_text = "、".join("/".join(group) for group in missing_groups)
        result = {"rule_id": rule_id, "item": item, "status": status, "severity": severity, "evidence": evidence, "suggestion": ""}
        if chapter:
            result["requirement_level"] = source_or_level
        else:
            result["source_level"] = source_or_level
        if status != "存在":
            result["suggestion"] = f"建议补充或明确：{missing_text}" if missing_text else f"建议完善{item}相关内容。"
        results.append(result)
    return results


def chapter_status(doc: dict, keywords: list[str]) -> tuple[str, list[dict]]:
    evidence = evidence_for(doc, keywords, limit=1000)
    if not evidence:
        return "缺失", []
    heading_evidence = []
    for paragraph in doc["paragraphs"]:
        exact = paragraph.text in keywords
        numbered = paragraph.heading_kind in {"chapter", "section"} or paragraph.style.lower().startswith("heading")
        keyword_match = any(key in paragraph.text for key in keywords)
        if exact or (numbered and keyword_match):
            heading_evidence.append({"location": _paragraph_location(doc, paragraph), "text": paragraph.text, "matched": [key for key in keywords if key in paragraph.text], "order": int(paragraph.id.split(":")[1])})
    if heading_evidence:
        return "存在", heading_evidence[:5] + [item for item in evidence if item not in heading_evidence][:5]
    return "部分", evidence[:5]
def apply_generic_rule(doc: dict, rule: dict) -> dict:
    """通用必备条款规则执行器：先判适用条件（触发器未命中→不适用并附理由），
    再按内容要素分组与 match_mode 判定存在/部分/缺失。"""
    rule_id = rule["id"]; item = rule["item"]; severity = rule["severity"]
    triggers = rule.get("applies_triggers")
    if triggers and not any(word in doc["all_text"] for word in triggers):
        return {"rule_id": rule_id, "item": item, "status": "不适用", "severity": severity, "evidence": [], "suggestion": "", "inapplicable_reason": rule.get("applies_when", "")}
    groups = rule["groups"]; mode = rule.get("match_mode", "any")
    all_evidence = []; hit_flags = []
    for group in groups:
        ev = evidence_for(doc, group)
        hit_flags.append(bool(ev))
        for e in ev:
            if e not in all_evidence:
                all_evidence.append(e)
    if mode == "any":
        status = "存在" if hit_flags[0] else "缺失"
    elif mode == "all":
        status = "存在" if all(hit_flags) else "部分" if any(hit_flags) else "缺失"
    elif mode == "priority":
        status = "存在" if hit_flags[0] else "部分" if any(hit_flags[1:]) else "缺失"
    strong = rule.get("strong")
    if strong and status == "存在" and not any(word in doc["all_text"] for word in strong):
        status = "部分"
    missing_groups = [group for group, hit in zip(groups, hit_flags) if not hit]
    missing_text = "、".join("/".join(group) for group in missing_groups)
    suggestion = ""
    if status == "缺失":
        suggestion = f"建议补充或完善{item}相关内容。" + (f"覆盖：{missing_text}" if missing_text else "")
    elif status == "部分":
        if mode == "all":
            suggestion = f"已涉及但环节不完整，建议补充：{missing_text}"
        elif mode == "priority" and not hit_flags[0]:
            suggestion = f"仅见{('、'.join(groups[i][0] for i in range(1, len(groups)) if hit_flags[i])) or '相关'}表述，未形成实质内容，建议补充：{missing_text}"
        else:
            suggestion = f"建议补充或完善：{missing_text}" if missing_text else f"建议完善{item}相关内容。"
    result = {"rule_id": rule_id, "item": item, "status": status, "severity": severity, "evidence": all_evidence[:12], "suggestion": suggestion}
    if status == "不适用":
        result["inapplicable_reason"] = rule.get("applies_when", "")
    return result

def check_chapters(doc: dict, context: dict) -> list[dict]:
    if context["primary_type"] == "employee_handbook":
        results = []
        for rule_id, item, keywords, level, severity in EMPLOYEE_HANDBOOK_CHAPTERS:
            status, evidence = chapter_status(doc, keywords)
            results.append({"rule_id": rule_id, "item": item, "status": status, "requirement_level": level, "severity": severity, "evidence": evidence, "suggestion": "建议补充或集中设置该章节/主题，覆盖：" + "、".join(keywords[:4]) if status != "存在" else ""})
        return results
    if context["primary_type"] == "company_charter":
        results = []
        for rule_id, item, keywords, level, severity in COMPANY_CHARTER_CHAPTERS:
            status, evidence = chapter_status(doc, keywords)
            results.append({"rule_id": rule_id, "item": item, "status": status, "requirement_level": level, "severity": severity, "evidence": evidence, "suggestion": "建议补充或集中设置该章节/主题，覆盖：" + "、".join(keywords[:4]) if status != "存在" else ""})
        return results
    if context["primary_type"] == "meeting_governance":
        return check_grouped_rules(doc, MEETING_CHAPTER_RULES, chapter=True)
    results = []
    for rule_id, item, keywords, level, severity in MANAGEMENT_CHAPTERS:
        status, evidence = chapter_status(doc, keywords)
        if rule_id == "S-MA-008" and status == "存在":
            coverage = sum(word in doc["all_text"] for word in ["解释", "施行", "生效", "执行"])
            if coverage < 2:
                status = "部分"
        results.append({"rule_id": rule_id, "item": item, "status": status, "requirement_level": level, "severity": severity, "evidence": evidence, "suggestion": "建议补充或完善该章节，覆盖：" + "、".join(keywords[:4]) if status not in {"存在", "不适用"} else ""})
    return results

def check_mandatory(doc: dict, context: dict) -> list[dict]:
    if context["primary_type"] == "employee_handbook":
        results = check_grouped_rules(doc, EMPLOYEE_HANDBOOK_RULES)
        special_training_or_noncompete = any(term in doc["all_text"] for term in ["专项培训费用", "竞业限制", "竞业禁止"]) or bool(re.search(r"服务期(?:限|协议|约定|为|不得)", doc["all_text"]))
        if not special_training_or_noncompete:
            for result in results:
                if result["rule_id"] == "M-LAB-011":
                    result.update({"status": "不适用", "evidence": [], "suggestion": ""})
        return results
    if context["primary_type"] == "company_charter":
        return check_grouped_rules(doc, COMPANY_CHARTER_RULES)
    if context["primary_type"] == "meeting_governance":
        return check_grouped_rules(doc, MEETING_RULES)
    if context["primary_type"] == "procurement_agent":
        rules = PROCUREMENT_AGENT_RULES
        results = []
        for rule_id, item, keywords, source, severity in rules:
            evidence = evidence_for(doc, keywords)
            status = "缺失" if not evidence else "存在" if len(evidence) >= 2 else "部分"
            results.append({"rule_id": rule_id, "item": item, "status": status, "source_level": source, "severity": severity, "evidence": evidence, "suggestion": "建议补充或完善该条款，覆盖：" + "、".join(keywords[:4]) if status != "存在" else ""})
        return results
    results = []
    for rule in GENERIC_RULES:
        results.append(apply_generic_rule(doc, rule))
    results.append({"rule_id": "TYPE-PENDING", "item": "专门类型必备条款清单", "status": "待确认", "severity": "中", "evidence": [], "suggestion": "当前清单未配置该制度类型的专门条款，请维护 references/mandatory_clause_catalog.md 和脚本规则数据结构。"})
    return results

def coverage(results: list[dict]) -> dict:
    ids = [item["rule_id"] for item in results]
    duplicates = sorted({rule_id for rule_id in ids if ids.count(rule_id) > 1})
    return {"selected_rules": len(ids), "executed_rules": len(ids), "coverage_rate": 1.0 if not duplicates else 0.0, "duplicate_rule_ids": duplicates, "missing_rule_ids": []}

def markdown_report(audits: list[dict]) -> str:
    lines = ["# 制度完整性审核报告", "", "> 审核基线：制度起草指引要求.md、不同制度类型必备条款清单.md", "> 审核范围：章节完整性和必备条款，不包含字体字号、页面排版和命名规范的完整审核。", ""]
    for audit in audits:
        all_results = audit["mandatory_clause_checks"] + audit["chapter_checks"]
        concerns = [item for item in all_results if item["status"] not in {"存在", "不适用"}]
        high_count = sum(item["severity"] == "高" and item["status"] in {"缺失", "部分"} for item in all_results)
        matched_orgs = "、".join(item["canonical_name"] for item in audit["classification"].get("matched_organizations", [])) or "未识别特定组织"
        lines += [f"## {audit['document']['file_name']}", "", f"- 制度名称：{audit['document']['title']}", f"- 主类型：{audit['classification']['primary_type']}（置信度：{audit['classification']['type_confidence']}）", f"- 名称层级：{audit['classification']['declared_level']}", f"- 组织范围：{audit['classification'].get('organization_scope', 'unresolved')}；命中组织：{matched_orgs}", f"- 需关注项：{len(concerns)}项，其中高风险 {high_count}项", f"- 审核执行覆盖率：{audit['coverage']['coverage_rate']:.0%}", ""]
        lines += ["### 必备条款审核", "", "| 编号 | 条款 | 状态 | 严重程度 | 证据/位置 | 建议 |", "|---|---|---|---|---|---|"]
        rows = [x for x in audit["mandatory_clause_checks"] if x["status"] not in {"存在", "不适用"}]
        rows.sort(key=concern_sort_key)
        if not rows:
            lines.append("| - | 无 | 已发现相关内容 | - | - | - |")
        else:
            for item in rows:
                evidence = "; ".join(x["location"] for x in item["evidence"]) or "未发现"
                lines.append(f"| {item['rule_id']} | {item['item']} | {item['status']} | {item['severity']} | {evidence} | {item['suggestion']} |")
        lines += ["", "### 章节审核", "", "| 编号 | 章节/主题 | 状态 | 严重程度 | 证据/位置 | 建议 |", "|---|---|---|---|---|---|"]
        rows = [x for x in audit["chapter_checks"] if x["status"] not in {"存在", "不适用"}]
        rows.sort(key=concern_sort_key)
        if not rows:
            lines.append("| - | 无 | 已发现相关内容 | - | - | - |")
        else:
            for item in rows:
                evidence = "; ".join(x["location"] for x in item["evidence"]) or "未发现"
                lines.append(f"| {item['rule_id']} | {item['item']} | {item['status']} | {item['severity']} | {evidence} | {item['suggestion']} |")
        lines += ["", "### 适用性说明", "", "- 位置一律使用章条坐标（如“第一章第三条”）；未发现对应内容时以“未发现”表示。", "- “不适用”表示根据制度调整事项未触发该审核点，不作为缺陷。", "- “待确认”表示证据不足、类型无法可靠识别或尚未配置专门规则，不等同于条款缺失。", ""]
    return "\n".join(lines)

def run(paths: list[Path], output_dir: Path) -> list[dict]:
    output_dir.mkdir(parents=True, exist_ok=True)
    audits = []
    for path in paths:
        doc = extract_docx(path)
        context = classify(doc)
        mandatory = check_mandatory(doc, context)
        chapters = check_chapters(doc, context)
        audit = {"audit_info": {"parsed_at": datetime.now().isoformat(timespec="seconds"), "basis": ["references/company_drafting_standard.md", "references/mandatory_clause_catalog.md", "references/organization_structure.json"], "profile": QUICK_AUDIT_PROFILE}, "document": {"file_name": doc["file_name"], "source_file": doc["path"], "title": doc["title"]}, "classification": context, "mandatory_clause_checks": mandatory, "chapter_checks": chapters, "coverage": coverage(mandatory + chapters), "parse_issues": doc["parse_issues"], "extracted": serialize_doc(doc)}
        audits.append(audit)
        safe = re.sub(r"[\\/:*?\"<>|]", "_", path.stem)
        (output_dir / f"{safe}.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    report_name = "制度完整性审核报告.md"
    summary_name = "制度完整性审核汇总.json"
    (output_dir / report_name).write_text(markdown_report(audits), encoding="utf-8")
    (output_dir / summary_name).write_text(json.dumps({"audits": audits}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"documents": len(audits), "output_dir": str(output_dir), "coverage": [audit["coverage"] for audit in audits]}, ensure_ascii=False, indent=2))
    return audits

def main() -> None:
    parser = argparse.ArgumentParser(description="按制度完整性清单审核 DOCX")
    parser.add_argument("documents", nargs="+", type=Path)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    args = parser.parse_args()
    run(args.documents, args.output_dir)

if __name__ == "__main__":
    main()



