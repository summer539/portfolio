from __future__ import annotations

import argparse
import json
import re
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from docx import Document

SKILL_ROOT = Path(__file__).resolve().parent.parent


# ---------------------------------------------------------------------------
# 1. DOCX extraction
# ---------------------------------------------------------------------------

def clean(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def heading_kind(text: str, style: str) -> str | None:
    if style.lower().startswith("heading"):
        return style
    if re.match(r"^第[一二三四五六七八九十百]+章", text):
        return "chapter"
    if re.match(r"^第[一二三四五六七八九十百]+条", text):
        return "article"
    return None


@dataclass
class Para:
    id: str
    text: str
    heading: str | None = None
    article_no: str | None = None


@dataclass
class DocInfo:
    path: str
    file_name: str
    title: str
    paragraphs: list[Para] = field(default_factory=list)
    all_text: str = ""
    parsed_issues: list[str] = field(default_factory=list)


def extract_docx(path: Path) -> DocInfo:
    doc = Document(str(path))
    paragraphs: list[Para] = []
    for idx, p in enumerate(doc.paragraphs, 1):
        text = clean(p.text)
        if not text:
            continue
        hk = heading_kind(text, p.style.name)
        art = None
        m = re.match(r"^第([一二三四五六七八九十百]+)条", text)
        if m:
            art = m.group(0)
        paragraphs.append(Para(f"para:{idx}", text, hk, art))
    tables: list[str] = []
    for table in doc.tables:
        rows = [" | ".join(clean(c.text) for c in row.cells) for row in table.rows]
        tables.append("\n".join(rows))
    all_text = "\n".join(p.text for p in paragraphs) + "\n" + "\n".join(tables)
    # detect title
    title = "未识别"
    for p in paragraphs[:10]:
        core = re.sub(r"\s+", "", p.text)
        for suffix in ["管理规定", "管理办法", "操作规程", "实施细则", "工作规则", "工作细则"]:
            if core.endswith(suffix) and not core.startswith("第"):
                title = core
                break
    return DocInfo(str(path), path.name, title, paragraphs, all_text)


# ---------------------------------------------------------------------------
# 2. Key information extraction
# ---------------------------------------------------------------------------

FEE_PATTERN = re.compile(r"(货物|服务|工程)采购费率为(\d+\.\d+)%")
FEE_PATTERN_ALT = re.compile(r"(货物|服务|工程).{0,6}费率为(\d+\.\d+)%")
THRESHOLD_PATTERN = re.compile(r"(\d+)万元(?:以上|以下)")
PENALTY_PATTERN = re.compile(r"扣罚代理服务费的(\d+)%")
CAP_PATTERN = re.compile(r"上限金额为(\d+)万")
PERIOD_PATTERN = re.compile(r"(年度|半年度|半年|季度|每半年|每季度|每年)")
SELECTION_PATTERN = re.compile(r"(公开寻源|邀请招标|直接采购|直接指定)")
APPROVAL_PATTERN = re.compile(r"(总部供应链管理部|省分公司供应链管理中心|采购实施部门|公司采购部|自行确定|自主选择)")
DEFINITION_PATTERN = re.compile(r"(?:所称|是指).{0,80}")


def extract_key_info(doc: DocInfo) -> dict:
    text = doc.all_text
    info: dict[str, list] = {}

    # Fee rates
    fees: list[dict] = []
    for pat in [FEE_PATTERN, FEE_PATTERN_ALT]:
        for m in pat.finditer(text):
            entry = {"category": m.group(1), "rate": m.group(2)}
            if entry not in fees:
                fees.append(entry)
    info["fee_rates"] = fees

    # Thresholds
    thresholds: list[dict] = []
    for p in doc.paragraphs:
        for m in THRESHOLD_PATTERN.finditer(p.text):
            direction = "以上" if "以上" in p.text[m.start():m.end()] else "以下"
            thresholds.append({"value": int(m.group(1)), "direction": direction, "location": p.id, "text": p.text[:120]})
    info["thresholds"] = thresholds

    # Penalties
    penalties: list[dict] = []
    for p in doc.paragraphs:
        for m in PENALTY_PATTERN.finditer(p.text):
            penalties.append({"rate": int(m.group(1)), "location": p.id, "text": p.text[:120]})
    info["penalties"] = penalties

    # Direct purchase cap
    caps: list[dict] = []
    for p in doc.paragraphs:
        for m in CAP_PATTERN.finditer(p.text):
            caps.append({"value": int(m.group(1)), "location": p.id, "text": p.text[:120]})
    info["caps"] = caps

    # Evaluation periods
    periods: list[dict] = []
    for p in doc.paragraphs:
        for m in PERIOD_PATTERN.finditer(p.text):
            if any(k in p.text for k in ["考核", "评估", "评价"]):
                periods.append({"period": m.group(1), "location": p.id, "text": p.text[:120]})
    info["evaluation_periods"] = periods

    # Selection methods
    methods: list[dict] = []
    for p in doc.paragraphs:
        for m in SELECTION_PATTERN.finditer(p.text):
            methods.append({"method": m.group(1), "location": p.id, "text": p.text[:120]})
    info["selection_methods"] = methods

    # Approval authorities
    approvals: list[dict] = []
    for p in doc.paragraphs:
        for m in APPROVAL_PATTERN.finditer(p.text):
            if any(k in p.text for k in ["审批", "备案", "确定", "选择", "自主", "自行"]):
                approvals.append({"entity": m.group(1), "location": p.id, "text": p.text[:120]})
    info["approval_entities"] = approvals

    # Term definitions
    definitions: list[dict] = []
    for p in doc.paragraphs:
        if "所称" in p.text or ("是指" in p.text and ("采购代理机构" in p.text or "采购" in p.text[:20])):
            definitions.append({"text": p.text[:200], "location": p.id})
    info["definitions"] = definitions

    return info


# ---------------------------------------------------------------------------
# 3. Internal consistency checks (single document)
# ---------------------------------------------------------------------------

@dataclass
class Conflict:
    id: str
    topic: str
    dimension: str
    type: str  # 硬冲突 / 软不一致
    severity: str
    evidence: list[dict]
    suggestion: str = ""


def check_internal(doc: DocInfo, info: dict) -> list[Conflict]:
    conflicts: list[Conflict] = []
    cid = 0

    def next_id():
        nonlocal cid
        cid += 1
        return f"IC-{cid}"

    # A: Term definitions - same term defined differently
    defs = info.get("definitions", [])
    if len(defs) >= 2:
        # Group by what term is being defined
        for i in range(len(defs)):
            for j in range(i + 1, len(defs)):
                d1, d2 = defs[i], defs[j]
                t1, t2 = d1["text"], d2["text"]
                if "采购代理机构" in t1 and "采购代理机构" in t2:
                    # Check if definitions differ
                    if "社会中介组织" in t1 and "企业法人" in t2:
                        conflicts.append(Conflict(
                            next_id(), "采购代理机构定义矛盾", "A", "硬冲突", "高",
                            [{"location": d1["location"], "excerpt": t1[:120]},
                             {"location": d2["location"], "excerpt": t2[:120]}],
                            "建议统一定义。'社会中介组织'与'企业法人'内涵不同"
                        ))
                    elif t1[:30] != t2[:30] and ("中介组织" in t1) != ("中介组织" in t2):
                        conflicts.append(Conflict(
                            next_id(), "采购代理机构定义不一致", "A", "硬冲突", "高",
                            [{"location": d1["location"], "excerpt": t1[:120]},
                             {"location": d2["location"], "excerpt": t2[:120]}],
                            "建议统一定义口径"
                        ))

    # B: Penalty standards - same violation different rates
    penalties = info.get("penalties", [])
    if len(penalties) >= 2:
        rates = [p["rate"] for p in penalties]
        if len(set(rates)) > 1:
            conflicts.append(Conflict(
                next_id(), "擅自更换人员处罚标准矛盾", "B", "硬冲突", "高",
                [{"location": p["location"], "excerpt": p["text"]} for p in penalties],
                f"建议统一处罚标准。同一违规行为罚率不同：{'、'.join(str(r) + '%' for r in rates)}"
            ))

    # C: Approval thresholds - same matter different thresholds
    thresholds = info.get("thresholds", [])
    hq_thresholds = []
    seen = set()
    for t in thresholds:
        if "总部" in t["text"] and "审批" in t["text"]:
            key = (t["value"], t["direction"])
            if key not in seen:
                seen.add(key)
                hq_thresholds.append(t)
    if len(hq_thresholds) >= 2:
        vals = [t["value"] for t in hq_thresholds]
        if len(set(vals)) > 1:
            conflicts.append(Conflict(
                next_id(), "审批权限阈值矛盾", "C", "硬冲突", "高",
                [{"location": t["location"], "excerpt": t["text"]} for t in hq_thresholds],
                f"建议统一审批阈值。同一事项规定不同：{'、'.join(str(v) + '万元' for v in vals)}"
            ))

    # E: Evaluation periods - different periods within same doc
    periods = info.get("evaluation_periods", [])
    if len(periods) >= 2:
        period_vals = set()
        seen_periods = []
        for p in periods:
            if "半年" in p["period"]:
                v = "半年"
            elif "季度" in p["period"]:
                v = "季度"
            elif "年度" in p["period"] or "每年" in p["period"]:
                v = "年度"
            else:
                continue
            period_vals.add(v)
            if not any(s["location"] == p["location"] for s in seen_periods):
                seen_periods.append(p)
        if len(period_vals) > 1:
            conflicts.append(Conflict(
                next_id(), "考核评估频次矛盾", "E", "硬冲突", "高",
                [{"location": p["location"], "excerpt": p["text"]} for p in seen_periods],
                f"建议统一考核频次。文档内出现多种频次：{'、'.join(period_vals)}"
            ))

    # F: Fee rates - different rates within same doc
    fees = info.get("fee_rates", [])
    if fees:
        cat_groups: dict[str, list[str]] = {}
        for f in fees:
            cat_groups.setdefault(f["category"], []).append(f["rate"])
        for cat, rates in cat_groups.items():
            if len(set(rates)) > 1:
                conflicts.append(Conflict(
                    next_id(), f"{cat}采购费率矛盾", "F", "硬冲突", "高",
                    [{"location": "doc", "excerpt": f"{cat}费率：{' vs '.join(rates)}"}],
                    f"建议统一{cat}采购费率。同一金额区间费率不同：{' vs '.join(rates)}"
                ))

    return conflicts


# ---------------------------------------------------------------------------
# 4. Cross-document consistency checks
# ---------------------------------------------------------------------------

def check_cross(docs: list[tuple[DocInfo, dict]]) -> list[Conflict]:
    conflicts: list[Conflict] = []
    cid = 0

    def next_id():
        nonlocal cid
        cid += 1
        return f"XC-{cid}"

    if len(docs) < 2:
        return conflicts

    # Compare every pair
    for i in range(len(docs)):
        for j in range(i + 1, len(docs)):
            d1, i1 = docs[i]
            d2, i2 = docs[j]
            n1, n2 = d1.file_name, d2.file_name

            # F: Fee rates
            fees1 = {(f["category"], f["rate"]) for f in i1.get("fee_rates", [])}
            fees2 = {(f["category"], f["rate"]) for f in i2.get("fee_rates", [])}
            common_cats = {f[0] for f in fees1} & {f[0] for f in fees2}
            for cat in common_cats:
                r1 = [f[1] for f in fees1 if f[0] == cat]
                r2 = [f[1] for f in fees2 if f[0] == cat]
                if r1 and r2 and r1[0] != r2[0]:
                    conflicts.append(Conflict(
                        next_id(), f"{cat}采购费率矛盾", "F", "硬冲突", "高",
                        [{"document": n1, "location": "doc", "excerpt": f"{cat}费率{r1[0]}%"},
                         {"document": n2, "location": "doc", "excerpt": f"{cat}费率{r2[0]}%"}],
                        f"建议统一{cat}采购费率（{r1[0]}% vs {r2[0]}%）"
                    ))

            # F: Direct purchase caps
            caps1 = i1.get("caps", [])
            caps2 = i2.get("caps", [])
            if caps1 and caps2:
                v1, v2 = caps1[0]["value"], caps2[0]["value"]
                if v1 != v2:
                    conflicts.append(Conflict(
                        next_id(), "直接采购费用上限矛盾", "F", "硬冲突", "高",
                        [{"document": n1, "location": caps1[0]["location"], "excerpt": caps1[0]["text"][:120]},
                         {"document": n2, "location": caps2[0]["location"], "excerpt": caps2[0]["text"][:120]}],
                        f"建议统一上限（{v1}万 vs {v2}万）"
                    ))

            # E: Evaluation periods
            p1 = i1.get("evaluation_periods", [])
            p2 = i2.get("evaluation_periods", [])
            if p1 and p2:
                pv1 = p1[0]["period"]
                pv2 = p2[0]["period"]
                norm1 = "半年" if "半年" in pv1 else "季度" if "季度" in pv1 else "年度" if "年" in pv1 else pv1
                norm2 = "半年" if "半年" in pv2 else "季度" if "季度" in pv2 else "年度" if "年" in pv2 else pv2
                if norm1 != norm2:
                    conflicts.append(Conflict(
                        next_id(), "考核频次矛盾", "E", "硬冲突", "高",
                        [{"document": n1, "location": p1[0]["location"], "excerpt": p1[0]["text"][:120]},
                         {"document": n2, "location": p2[0]["location"], "excerpt": p2[0]["text"][:120]}],
                        f"建议统一考核频次（{norm1} vs {norm2}）"
                    ))

            # C: Selection methods - 邀请招标 conflict
            m1 = {m["method"] for m in i1.get("selection_methods", [])}
            m2 = {m["method"] for m in i2.get("selection_methods", [])}
            if "邀请招标" in m1 and "邀请招标" not in m2 and "公开寻源" in m2:
                conflicts.append(Conflict(
                    next_id(), "选择方式矛盾", "C", "硬冲突", "高",
                    [{"document": n1, "location": "doc", "excerpt": "允许邀请招标"},
                     {"document": n2, "location": "doc", "excerpt": "不得邀请招标/仅公开寻源"}],
                    f"建议统一选择方式，{n1}允许邀请招标但{n2}不允许"
                ))
            elif "邀请招标" in m2 and "邀请招标" not in m1 and "公开寻源" in m1:
                conflicts.append(Conflict(
                    next_id(), "选择方式矛盾", "C", "硬冲突", "高",
                    [{"document": n1, "location": "doc", "excerpt": "仅公开寻源"},
                     {"document": n2, "location": "doc", "excerpt": "允许邀请招标"}],
                    f"建议统一选择方式，{n2}允许邀请招标但{n1}不允许"
                ))

            # C/D: Approval authority - self-selection vs unified approval
            a1 = {a["entity"] for a in i1.get("approval_entities", [])}
            a2 = {a["entity"] for a in i2.get("approval_entities", [])}
            if ("自行确定" in a1 or "自主选择" in a1) and "总部供应链管理部" in a2:
                conflicts.append(Conflict(
                    next_id(), "审批权限矛盾", "C", "硬冲突", "高",
                    [{"document": n1, "location": "doc", "excerpt": "自主选择/自行确定"},
                     {"document": n2, "location": "doc", "excerpt": "总部统一审批"}],
                    f"建议统一审批权限，{n1}主张自主选择但{n2}要求总部统一审批"
                ))
            elif ("自行确定" in a2 or "自主选择" in a2) and "总部供应链管理部" in a1:
                conflicts.append(Conflict(
                    next_id(), "审批权限矛盾", "C", "硬冲突", "高",
                    [{"document": n1, "location": "doc", "excerpt": "总部统一审批"},
                     {"document": n2, "location": "doc", "excerpt": "自主选择/自行确定"}],
                    f"建议统一审批权限，{n2}主张自主选择但{n1}要求总部统一审批"
                ))

            # B: Penalty standards
            pen1 = i1.get("penalties", [])
            pen2 = i2.get("penalties", [])
            if pen1 and pen2:
                r1, r2 = pen1[0]["rate"], pen2[0]["rate"]
                if r1 != r2:
                    conflicts.append(Conflict(
                        next_id(), "处罚标准矛盾", "B", "硬冲突", "高",
                        [{"document": n1, "location": pen1[0]["location"], "excerpt": pen1[0]["text"][:120]},
                         {"document": n2, "location": pen2[0]["location"], "excerpt": pen2[0]["text"][:120]}],
                        f"建议统一处罚标准（{r1}% vs {r2}%）"
                    ))

            # G: Negative behavior management
            if "不适用" in d1.all_text and "负面行为管理规范" in d2.all_text:
                conflicts.append(Conflict(
                    next_id(), "负面行为管理规范矛盾", "G", "硬冲突", "高",
                    [{"document": n1, "location": "doc", "excerpt": "自行制定/不适用总部规范"},
                     {"document": n2, "location": "doc", "excerpt": "纳入总部负面行为管理规范"}],
                    "建议明确子公司与总部负面行为管理规范的衔接关系"
                ))

            # D: Department naming
            if "供应链管理部" in d1.all_text and "公司采购部" in d2.all_text and "供应链管理部" not in d2.all_text:
                conflicts.append(Conflict(
                    next_id(), "归口部门称谓不一致", "D", "软不一致", "中",
                    [{"document": n1, "location": "doc", "excerpt": "供应链管理部"},
                     {"document": n2, "location": "doc", "excerpt": "公司采购部"}],
                    "建议统一归口部门名称或明确对应关系"
                ))

    return conflicts


# ---------------------------------------------------------------------------
# 5. Report generation
# ---------------------------------------------------------------------------

def make_report(docs: list[tuple[DocInfo, dict]], internal: list[Conflict],
                cross: list[Conflict], elapsed: float) -> str:
    lines = [
        "# 制度一致性快速审核报告", "",
        f"> 审核基线：制度一致性审核清单.md（V1.0）",
        f"> 审核日期：{datetime.now().strftime('%Y-%m-%d')}",
        f"> 审核耗时：{elapsed:.2f}秒",
        f"> 审核对象：{len(docs)}份文件", ""
    ]

    # Overview
    all_conflicts = internal + cross
    hard = [c for c in all_conflicts if c.type == "硬冲突"]
    soft = [c for c in all_conflicts if c.type == "软不一致"]
    lines += [
        "## 审核概览", "",
        f"- 文件数：{len(docs)}",
        f"- 内部冲突：{len(internal)}项",
        f"- 跨制度冲突：{len(cross)}项",
        f"- 硬冲突：{len(hard)}项",
        f"- 软不一致：{len(soft)}项", ""
    ]

    # Document list
    lines += ["## 审核文件清单", "", "| 序号 | 文件名 | 制度名称 | 条款数 |", "|---|---|---|---|"]
    for idx, (doc, _) in enumerate(docs, 1):
        lines.append(f"| {idx} | {doc.file_name} | {doc.title} | {len(doc.paragraphs)} |")
    lines.append("")

    # Conflict list
    lines += ["## 冲突清单（按严重度降序）", ""]
    sorted_c = sorted(all_conflicts, key=lambda c: {"高": 0, "中": 1, "低": 2}[c.severity])

    if not sorted_c:
        lines += ["未发现冲突项。", ""]
    else:
        lines += ["| 编号 | 冲突点 | 维度 | 类型 | 严重度 | 证据 | 修改建议 |", "|---|---|---|---|---|---|---|"]
        for c in sorted_c:
            ev = "; ".join(f"[{e.get('document', e.get('location', ''))}] {e.get('excerpt', '')[:60]}" for e in c.evidence)
            lines.append(f"| {c.id} | {c.topic} | {c.dimension} | {c.type} | {c.severity} | {ev} | {c.suggestion} |")
    lines.append("")

    # Term matrix
    all_defs: dict[str, list] = {}
    for doc, info in docs:
        for d in info.get("definitions", []):
            all_defs.setdefault("采购代理机构", []).append({"doc": doc.file_name, "text": d["text"][:100]})
    if all_defs:
        lines += ["## 术语一致性矩阵", "", "| 术语 | 定义1 | 定义2 | 是否一致 | 说明 |", "|---|---|---|---|---|"]
        for term, defs in all_defs.items():
            if len(defs) >= 2:
                d1 = defs[0]["text"][:50]
                d2 = defs[1]["text"][:50] if len(defs) > 1 else ""
                consistent = "一致" if d1[:20] == d2[:20] else "不一致"
                lines.append(f"| {term} | {d1} | {d2} | {consistent} | — |")
        lines.append("")

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# 6. Main
# ---------------------------------------------------------------------------

def safe_name(value: str) -> str:
    return re.sub(r'[\\/:*?"<>|]', "_", value)


def run(documents: list[Path], output_dir: Path) -> dict:
    start = time.perf_counter()
    output_dir.mkdir(parents=True, exist_ok=True)

    docs: list[tuple[DocInfo, dict]] = []
    for path in documents:
        doc = extract_docx(path)
        info = extract_key_info(doc)
        docs.append((doc, info))

    # Internal consistency (each doc individually)
    all_internal: list[Conflict] = []
    for doc, info in docs:
        conflicts = check_internal(doc, info)
        for c in conflicts:
            c.id = f"{doc.file_name[:6]}-{c.id}"
        all_internal.extend(conflicts)

    # Cross-document consistency
    all_cross = check_cross(docs)

    elapsed = time.perf_counter() - start

    # Generate report
    report = make_report(docs, all_internal, all_cross, elapsed)
    (output_dir / "一致性快速审核报告.md").write_text(report, encoding="utf-8")

    # Generate JSON
    result = {
        "audit_info": {
            "parsed_at": datetime.now().isoformat(timespec="seconds"),
            "elapsed_seconds": round(elapsed, 3),
            "documents": [{"file_name": d.file_name, "title": d.title} for d, _ in docs],
            "total_conflicts": len(all_internal) + len(all_cross),
            "hard_conflicts": len([c for c in all_internal + all_cross if c.type == "硬冲突"]),
            "soft_inconsistencies": len([c for c in all_internal + all_cross if c.type == "软不一致"]),
        },
        "internal_conflicts": [
            {"id": c.id, "topic": c.topic, "dimension": c.dimension, "type": c.type,
             "severity": c.severity, "evidence": c.evidence, "suggestion": c.suggestion}
            for c in all_internal
        ],
        "cross_conflicts": [
            {"id": c.id, "topic": c.topic, "dimension": c.dimension, "type": c.type,
             "severity": c.severity, "evidence": c.evidence, "suggestion": c.suggestion}
            for c in all_cross
        ],
    }
    (output_dir / "一致性快速审核汇总.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    summary = {
        "documents": len(docs),
        "internal_conflicts": len(all_internal),
        "cross_conflicts": len(all_cross),
        "hard_conflicts": result["audit_info"]["hard_conflicts"],
        "soft_inconsistencies": result["audit_info"]["soft_inconsistencies"],
        "elapsed_seconds": round(elapsed, 3),
        "output_dir": str(output_dir),
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="制度一致性快速审核")
    parser.add_argument("documents", nargs="+", type=Path, help="待审核的DOCX文件路径（至少1份）")
    parser.add_argument("--output-dir", type=Path, default=None, help="输出目录")
    args = parser.parse_args()

    default_name = safe_name(args.documents[0].stem) if len(args.documents) == 1 else "批量一致性审核"
    out = args.output_dir or (Path.cwd() / "审核输出" / "consistency" / default_name)
    run(args.documents, out)


if __name__ == "__main__":
    main()
